"""Java Anonymizer exported patient index (XLSX) I/O.

Public API
----------
JavaAnonymizerExportedStudy
    One row from the Java Anonymizer patient index workbook.

read_java_anonymizer_index_xlsx(filename: str) -> list[JavaAnonymizerExportedStudy]
    Read the Java Anonymizer exported patient index workbook.
"""
from dataclasses import dataclass
from typing import Union

from openpyxl import Workbook, load_workbook
from openpyxl.workbook.child import _WorkbookChild
from openpyxl.worksheet.worksheet import Worksheet


@dataclass
class JavaAnonymizerExportedStudy:
    """One row from the Java Anonymizer patient index workbook."""

    ANON_PatientName: str
    ANON_PatientID: str
    PHI_PatientName: str
    PHI_PatientID: str
    DateOffset: str
    ANON_StudyDate: str
    PHI_StudyDate: str
    ANON_Accession: str
    PHI_Accession: str
    ANON_StudyInstanceUID: str
    PHI_StudyInstanceUID: str


def read_java_anonymizer_index_xlsx(filename: str) -> list[JavaAnonymizerExportedStudy]:
    """
    Read the Java Anonymizer exported patient index workbook.

    Args:
        filename: Path to the Excel file.

    Returns:
        Rows as JavaAnonymizerExportedStudy instances (empty if the sheet has no data).

    Raises:
        ValueError: If no active worksheet is found.
        FileNotFoundError: If the file is missing.
    """
    workbook: Workbook = load_workbook(filename)
    sheet: Union[_WorkbookChild, None] = workbook.active
    data: list[JavaAnonymizerExportedStudy] = []

    if sheet is None or not isinstance(sheet, Worksheet):
        raise ValueError("No active sheet found in the workbook")

    for row in sheet.iter_rows(values_only=True, min_row=2):
        str_row = [str(item) if item is not None else "" for item in row]
        data.append(JavaAnonymizerExportedStudy(*str_row))

    return data
